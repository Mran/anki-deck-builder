#!/usr/bin/env python3
"""Convert learning materials into markdown under <project>/sources/, images into <project>/media/.

  ingest.py <file-or-url> [...] --project <dir> [--lang zh] [--whisper-model small]

Each input is handled independently; a missing optional dependency skips that input with an
install hint instead of aborting the rest. See references/ingestion.md for fallbacks.
"""
import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.request
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import dump_yaml, load_yaml  # noqa: E402

VIDEO = {".mp4", ".mkv", ".mov", ".webm", ".avi", ".flv"}
AUDIO = {".mp3", ".m4a", ".wav", ".aac", ".ogg", ".flac"}
IMAGE = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
VIDEO_HOSTS = ("youtube.com", "youtu.be", "bilibili.com", "b23.tv", "vimeo.com")


class Skip(Exception):
    pass


def need(module, pip_name=None):
    try:
        return __import__(module)
    except ImportError:
        raise Skip(f"需要安装依赖:pip install {pip_name or module}")


def slugify(s):
    s = re.sub(r"[\\/:*?\"<>|#\s]+", "-", s).strip("-.")
    return s[:60] or "source"


class Ctx:
    def __init__(self, project, lang, whisper_model):
        self.root = Path(project).resolve()
        self.media = self.root / "media"
        self.sources = self.root / "sources"
        self.media.mkdir(parents=True, exist_ok=True)
        self.sources.mkdir(parents=True, exist_ok=True)
        self.lang = lang
        self.whisper_model = whisper_model

    def save_image(self, data, stem, ext):
        name = f"{stem}{ext.lower()}"
        (self.media / name).write_bytes(data)
        return name


# ---------------- handlers: each returns (markdown, kind) ----------------

def h_text(path, ctx, stem):
    return path.read_text(encoding="utf-8", errors="replace"), "text"


def h_pdf(path, ctx, stem):
    pypdf = need("pypdf")
    reader = pypdf.PdfReader(str(path))
    out = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        out.append(f"## 第 {i} 页\n\n{text}\n")
    md = "\n".join(out)
    try:
        import fitz  # pymupdf
        doc = fitz.open(str(path))
        n = 0
        for pno, page in enumerate(doc, 1):
            for k, img in enumerate(page.get_images(full=True), 1):
                pix = fitz.Pixmap(doc, img[0])
                if pix.width < 120 or pix.height < 120:
                    continue
                if pix.n - pix.alpha >= 4:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                name = f"{stem}_p{pno}_{k}.png"
                pix.save(str(ctx.media / name))
                n += 1
        if n:
            md += f"\n\n<!-- 已提取 {n} 张图片到 media/{stem}_p*.png -->\n"
    except ImportError:
        pass
    if len(re.sub(r"\s|## 第 \d+ 页", "", md)) < 50 * max(1, len(reader.pages)) / 10:
        md = "<!-- 警告:提取到的文字极少,可能是扫描版 PDF,需要 OCR 或视觉读取 -->\n\n" + md
    return md, "pdf"


def h_docx(path, ctx, stem):
    docx = need("docx", "python-docx")
    d = docx.Document(str(path))
    out = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.match(r"(heading|标题)\s*(\d)", style)
            if m:
                out.append("#" * min(int(m.group(2)) + 1, 6) + " " + text)
            elif "title" in style:
                out.append("# " + text)
            elif "list" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            t = docx.table.Table(child, d)
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        out.append("")
    n = 0
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if info.filename.startswith("word/media/") and Path(info.filename).suffix.lower() in IMAGE:
                n += 1
                ctx.save_image(z.read(info), f"{stem}_{n}", Path(info.filename).suffix)
    if n:
        out.append(f"<!-- 已提取 {n} 张图片到 media/{stem}_*.* -->")
    return "\n".join(out), "docx"


def h_pptx(path, ctx, stem):
    pptx = need("pptx", "python-pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = pptx.Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## 第 {i} 页" + (f":{title}" if title else ""))
        img_n = 0

        def walk(shapes):
            nonlocal img_n
            for sh in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                    continue
                if slide.shapes.title is not None and sh.shape_id == slide.shapes.title.shape_id:
                    continue
                if sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        t = "".join(r.text for r in para.runs).strip()
                        if t:
                            out.append("  " * para.level + "- " + t)
                if getattr(sh, "has_table", False) and sh.has_table:
                    rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                    if rows:
                        out.append("| " + " | ".join(rows[0]) + " |")
                        out.append("|" + "---|" * len(rows[0]))
                        out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    img_n += 1
                    name = ctx.save_image(sh.image.blob, f"{stem}_s{i}_{img_n}", "." + sh.image.ext)
                    out.append(f"![]({name})")
        walk(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"\n> 演讲者备注:{notes}")
        out.append("")
    return "\n".join(out), "pptx"


def html_to_md(raw):
    try:
        import trafilatura
        md = trafilatura.extract(raw, output_format="markdown", include_tables=True, include_images=False)
        if md and len(md) > 200:
            return md
    except ImportError:
        pass
    bs4 = need("bs4", "beautifulsoup4")
    soup = bs4.BeautifulSoup(raw, "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript", "svg"]):
        t.decompose()
    main = soup.find("article") or soup.find("main") or soup.body or soup
    out = []
    for el in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote", "tr"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            out.append("\n" + "#" * int(el.name[1]) + " " + text)
        elif el.name == "li":
            out.append("- " + text)
        elif el.name == "pre":
            out.append("```\n" + el.get_text() + "\n```")
        elif el.name == "tr":
            out.append("| " + " | ".join(c.get_text(" ", strip=True) for c in el.find_all(["td", "th"])) + " |")
        else:
            out.append(text)
    return "\n".join(out)


def h_html(path, ctx, stem):
    return html_to_md(path.read_text(encoding="utf-8", errors="replace")), "html"


def h_url(url, ctx, stem):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (ankiskill ingest)"})
    with urllib.request.urlopen(req, timeout=30) as r:
        raw = r.read().decode(r.headers.get_content_charset() or "utf-8", errors="replace")
    md = html_to_md(raw)
    if len(md) < 200:
        raise Skip("抓取到的正文过少(可能需要登录或 JS 渲染),请改用 WebFetch 或浏览器工具读取")
    return md, "web"


def h_epub(path, ctx, stem):
    out = []
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))
        for n in names:
            md = html_to_md(z.read(n).decode("utf-8", errors="replace"))
            if md.strip():
                out.append(md)
    return "\n\n".join(out), "epub"


def parse_subs(text):
    text = text.lstrip("﻿")
    blocks = re.split(r"\n\s*\n", text.replace("\r", ""))
    out, last_mark, para = [], -999, []
    for b in blocks:
        lines = [l for l in b.split("\n") if l.strip()]
        ts = next((l for l in lines if "-->" in l), None)
        if not ts:
            continue
        m = re.match(r"\s*(?:(\d+):)?(\d+):(\d+)", ts)
        sec = (int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + int(m.group(3))) if m else 0
        content = " ".join(re.sub(r"<[^>]+>", "", l).strip() for l in lines[lines.index(ts) + 1:])
        if not content or (para and para[-1] == content):
            continue
        if sec - last_mark >= 60:
            if para:
                out.append(" ".join(para))
                para = []
            out.append(f"\n[{sec // 60:02d}:{sec % 60:02d}]")
            last_mark = sec
        para.append(content)
    if para:
        out.append(" ".join(para))
    return "\n".join(out)


def h_subs(path, ctx, stem):
    return parse_subs(path.read_text(encoding="utf-8", errors="replace")), "subtitles"


def transcribe(audio_path, ctx):
    try:
        from faster_whisper import WhisperModel
        model = WhisperModel(ctx.whisper_model, compute_type="int8")
        segs, _ = model.transcribe(str(audio_path), language=ctx.lang or None, vad_filter=True)
        segs = [(s.start, s.text) for s in segs]
    except ImportError:
        try:
            import whisper
            res = whisper.load_model(ctx.whisper_model).transcribe(str(audio_path), language=ctx.lang or None)
            segs = [(s["start"], s["text"]) for s in res["segments"]]
        except ImportError:
            raise Skip("本地转写需要:pip install faster-whisper(以及系统已安装 ffmpeg)")
    out, last, para = [], -999, []
    for start, text in segs:
        if start - last >= 60:
            if para:
                out.append("".join(para) if ctx.lang == "zh" else " ".join(para))
                para = []
            out.append(f"\n[{int(start) // 60:02d}:{int(start) % 60:02d}]")
            last = start
        para.append(text.strip())
    if para:
        out.append("".join(para) if ctx.lang == "zh" else " ".join(para))
    return "\n".join(out)


def h_media(path, ctx, stem):
    if not shutil.which("ffmpeg"):
        raise Skip("需要安装 ffmpeg 才能处理音视频")
    with tempfile.TemporaryDirectory() as td:
        wav = Path(td) / "audio.wav"
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), "-ac", "1", "-ar", "16000", str(wav)],
                       check=True)
        return transcribe(wav, ctx), "transcript"


def h_video_url(url, ctx, stem):
    if not shutil.which("yt-dlp"):
        raise Skip("需要安装 yt-dlp:pip install yt-dlp")
    langs = {"zh": "zh-Hans,zh-CN,zh,zh-Hant,ai-zh,en", "en": "en,en-US,en-GB"}.get(ctx.lang, "en")
    with tempfile.TemporaryDirectory() as td:
        base = [ "yt-dlp", "--skip-download", "--sub-langs", langs, "--sub-format", "vtt/srt/best",
                "--convert-subs", "srt", "-o", str(Path(td) / "v.%(ext)s"), url]
        subprocess.run(base[:2] + ["--write-subs"] + base[2:], capture_output=True)
        subs = sorted(Path(td).glob("*.srt"))
        if not subs:
            subprocess.run(base[:2] + ["--write-auto-subs"] + base[2:], capture_output=True)
            subs = sorted(Path(td).glob("*.srt"))
        if subs:
            return parse_subs(subs[0].read_text(encoding="utf-8", errors="replace")), "subtitles"
        print("  没有可用字幕,下载音频进行转写…")
        subprocess.run(["yt-dlp", "-x", "--audio-format", "m4a", "-o", str(Path(td) / "a.%(ext)s"), url],
                       check=True, capture_output=True)
        audio = next(Path(td).glob("a.*"), None)
        if not audio:
            raise Skip("音频下载失败")
        return h_media(audio, ctx, stem)


def h_image(path, ctx, stem):
    name = f"{stem}{path.suffix.lower()}"
    shutil.copy(path, ctx.media / name)
    return (f"![]({name})\n\n<!-- 图片材料:请用视觉能力读取 media/{name},把内容写在这里 -->\n"), "image"


def pick(src):
    if re.match(r"^https?://", src):
        return (h_video_url if any(h in src for h in VIDEO_HOSTS) else h_url), src
    p = Path(src)
    if not p.exists():
        raise Skip("文件不存在")
    ext = p.suffix.lower()
    table = {".txt": h_text, ".md": h_text, ".markdown": h_text, ".pdf": h_pdf, ".docx": h_docx,
             ".pptx": h_pptx, ".html": h_html, ".htm": h_html, ".epub": h_epub, ".srt": h_subs, ".vtt": h_subs}
    if ext in table:
        return table[ext], p
    if ext in VIDEO | AUDIO:
        return h_media, p
    if ext in IMAGE:
        return h_image, p
    if ext in (".doc", ".ppt"):
        raise Skip("旧版 Office 格式,请先另存为 .docx/.pptx(或用 LibreOffice: soffice --headless --convert-to docx)")
    raise Skip(f"不支持的格式 {ext}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--project", required=True)
    ap.add_argument("--lang", default=None, help="材料语言(影响字幕选择与转写),默认取 project.yaml 的 lang")
    ap.add_argument("--whisper-model", default="small")
    a = ap.parse_args()

    meta = load_yaml(Path(a.project) / "project.yaml") or sys.exit("不是项目目录,先运行 project.py init")
    ctx = Ctx(a.project, a.lang or meta.get("lang", "zh"), a.whisper_model)
    meta.setdefault("sources", [])
    ok = 0
    for src in a.inputs:
        try:
            handler, target = pick(src)
            stem = slugify(Path(str(target)).stem if not isinstance(target, str) else
                           re.sub(r"^https?://(www\.)?", "", target)[:60])
            dest = ctx.sources / f"{stem}.md"
            n = 2
            while dest.exists():
                dest = ctx.sources / f"{stem}-{n}.md"
                n += 1
            stem = dest.stem
            md, kind = handler(target, ctx, stem)
            header = f"---\norigin: {src}\nkind: {kind}\ningested: {dt.date.today().isoformat()}\n---\n\n"
            dest.write_text(header + md.strip() + "\n", encoding="utf-8")
            meta["sources"].append({"path": f"sources/{dest.name}", "origin": str(src), "kind": kind,
                                    "added": dt.date.today().isoformat()})
            print(f"✓ {src} → sources/{dest.name}({len(md)} 字符)")
            ok += 1
        except Skip as e:
            print(f"✗ {src}:{e}")
        except Exception as e:  # keep going with other inputs
            print(f"✗ {src}:{type(e).__name__}: {e}")
    dump_yaml(meta, Path(a.project) / "project.yaml")
    print(f"完成 {ok}/{len(a.inputs)}")


if __name__ == "__main__":
    main()
